from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .models import Resume
from .optimizer import NORMAL_PROFILE, LayoutProfile


FONT_NAME = "Arial"

PAGE_MARGIN = Inches(0.65)

NAME_SIZE = Pt(21)
HEADLINE_SIZE = Pt(11.5)
CONTACT_SIZE = Pt(9.5)
META_SIZE = Pt(9.5)
SUBTITLE_SIZE = Pt(10)

TAB_POSITION = Inches(7.2)

# Resume visual language: restrained corporate blue with neutral body text.
NAVY = RGBColor(0x17, 0x36, 0x5D)
BLUE = RGBColor(0x1F, 0x5E, 0x9C)
MUTED = RGBColor(0x66, 0x70, 0x85)
BLACK = RGBColor(0x20, 0x20, 0x20)


def configure_document(
    document: Document,
    profile: LayoutProfile = NORMAL_PROFILE,
) -> None:
    section = document.sections[0]

    section.top_margin = PAGE_MARGIN
    section.bottom_margin = PAGE_MARGIN
    section.left_margin = PAGE_MARGIN
    section.right_margin = PAGE_MARGIN

    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(profile.body_font_pt)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(
        profile.body_spacing_after_pt
    )
    normal.paragraph_format.line_spacing = 1.0

    if "Resume Bullet" not in styles:
        bullet_style = styles.add_style(
            "Resume Bullet",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    else:
        bullet_style = styles["Resume Bullet"]

    bullet_style.base_style = normal
    bullet_style.font.name = FONT_NAME
    bullet_style.font.size = Pt(profile.body_font_pt)
    bullet_style.font.color.rgb = BLACK
    bullet_style.paragraph_format.left_indent = Inches(0.22)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.14)
    # Slightly more breathing room than the global body rhythm, while
    # remaining conservative enough for the compact profiles.
    bullet_style.paragraph_format.space_after = Pt(
        max(profile.bullet_spacing_after_pt, 2.5)
    )
    bullet_style.paragraph_format.line_spacing = 1.0


def add_run(
    paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    size: Pt | None = None,
    color: RGBColor | None = None,
) -> None:
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color or BLACK

    if size is not None:
        run.font.size = size


def add_right_tab(paragraph) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        TAB_POSITION,
        WD_TAB_ALIGNMENT.RIGHT,
    )


def add_bottom_border(
    paragraph,
    *,
    color: str = "1F5E9C",
    size: str = "6",
) -> None:
    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)


def add_left_accent(
    paragraph,
    *,
    color: str = "1F5E9C",
    size: str = "14",
) -> None:
    """Add a subtle vertical accent without a full-width decorative rule."""
    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)

    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)

    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)


def add_header(document: Document, resume: Resume) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)

    add_run(
        paragraph,
        resume.contact.name,
        bold=True,
        size=NAME_SIZE,
        color=NAVY,
    )

    contact_parts = []

    if resume.contact.email:
        contact_parts.append(resume.contact.email)

    if resume.contact.phone:
        contact_parts.append(resume.contact.phone)

    if resume.contact.location:
        contact_parts.append(resume.contact.location)

    contact_parts.extend(resume.contact.links)

    if contact_parts:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)

        add_run(
            paragraph,
            " | ".join(contact_parts),
            size=CONTACT_SIZE,
            color=MUTED,
        )


def add_headline(document: Document, headline: str) -> None:
    if not headline:
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(7)

    add_run(
        paragraph,
        headline.upper(),
        bold=True,
        size=HEADLINE_SIZE,
        color=BLUE,
    )

    add_bottom_border(paragraph, color="1F5E9C", size="6")


def add_section_heading(
    document: Document,
    title: str,
    profile: LayoutProfile,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(
        profile.section_spacing_before_pt + 1
    )
    paragraph.paragraph_format.space_after = Pt(
        profile.section_spacing_after_pt + 1
    )
    paragraph.paragraph_format.keep_with_next = True

    add_run(
        paragraph,
        title.upper(),
        bold=True,
        size=Pt(profile.section_font_pt),
        color=BLUE,
    )

    add_left_accent(paragraph)


def add_summary(
    document: Document,
    summary: str,
    profile: LayoutProfile,
) -> None:
    if not summary:
        return

    add_section_heading(document, "Summary", profile)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(
        max(profile.body_spacing_after_pt, 3.5)
    )

    add_run(
        paragraph,
        summary,
        size=Pt(profile.body_font_pt),
    )


def add_skills(
    document: Document,
    resume: Resume,
    profile: LayoutProfile,
) -> None:
    if not resume.skills:
        return

    add_section_heading(document, "Core Skills", profile)

    for category in resume.skills:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(
            max(profile.body_spacing_after_pt, 2.5)
        )

        if category.name:
            add_run(
                paragraph,
                f"{category.name}: ",
                bold=True,
                size=Pt(profile.body_font_pt),
                color=NAVY,
            )

        add_run(
            paragraph,
            ", ".join(category.skills),
            size=Pt(profile.body_font_pt),
        )


def add_education(
    document: Document,
    resume: Resume,
    profile: LayoutProfile,
) -> None:
    if not resume.education:
        return

    add_section_heading(document, "Education", profile)

    for education in resume.education:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        add_right_tab(paragraph)

        add_run(
            paragraph,
            education.degree or education.institution,
            bold=True,
            size=Pt(profile.body_font_pt),
            color=NAVY,
        )

        if education.dates:
            add_run(
                paragraph,
                f"\t{education.dates}",
                size=META_SIZE,
                color=MUTED,
            )

        if education.degree:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(
                max(profile.body_spacing_after_pt, 2.5)
            )

            add_run(
                paragraph,
                education.institution,
                size=Pt(profile.body_font_pt),
            )

        for detail in education.details:
            bullet = document.add_paragraph(style="Resume Bullet")
            add_run(
                bullet,
                f"• {detail}",
                size=Pt(profile.body_font_pt),
            )


def add_certifications(
    document: Document,
    resume: Resume,
    profile: LayoutProfile,
) -> None:
    if not resume.certifications:
        return

    add_section_heading(document, "Certifications", profile)

    for certification in resume.certifications:
        paragraph = document.add_paragraph()
        # Certifications remain intentionally tighter than narrative
        # sections so Page 1 gains breathing room without losing content.
        paragraph.paragraph_format.space_after = Pt(
            min(profile.body_spacing_after_pt, 1.0)
        )
        add_right_tab(paragraph)

        add_run(
            paragraph,
            certification.name,
            size=Pt(profile.body_font_pt),
        )

        if certification.issuer:
            add_run(
                paragraph,
                f" | {certification.issuer}",
                size=Pt(profile.body_font_pt),
            )

        if certification.date:
            add_run(
                paragraph,
                f"\t{certification.date}",
                size=META_SIZE,
                color=MUTED,
            )


def add_experience(
    document: Document,
    resume: Resume,
    profile: LayoutProfile,
) -> None:
    if not resume.experience:
        return

    add_section_heading(
        document,
        "Professional Experience",
        profile,
    )

    for experience_index, experience in enumerate(resume.experience):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        add_right_tab(paragraph)

        add_run(
            paragraph,
            experience.company,
            bold=True,
            size=Pt(profile.body_font_pt + 0.5),
            color=NAVY,
        )

        if experience.location:
            add_run(
                paragraph,
                f"\t{experience.location}",
                size=META_SIZE,
                color=MUTED,
            )

        if experience.subtitle:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True

            add_run(
                paragraph,
                experience.subtitle,
                italic=True,
                size=SUBTITLE_SIZE,
                color=MUTED,
            )

        paragraph = document.add_paragraph()
        # Give the role line enough separation to read as a distinct
        # hierarchy level instead of another line in the company block.
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_with_next = True
        add_right_tab(paragraph)

        add_run(
            paragraph,
            experience.title,
            bold=True,
            size=Pt(profile.body_font_pt),
            color=BLUE,
        )

        if experience.dates:
            add_run(
                paragraph,
                f"\t{experience.dates}",
                size=META_SIZE,
                color=MUTED,
            )

        for index, bullet in enumerate(experience.bullets):
            paragraph = document.add_paragraph(style="Resume Bullet")
            is_last = index == len(experience.bullets) - 1

            if is_last:
                # Deliberately create visual separation between positions.
                # This is the main page-balance adjustment.
                spacing = max(profile.entry_spacing_after_pt + 3, 6)
                if experience_index == len(resume.experience) - 1:
                    spacing = max(profile.entry_spacing_after_pt + 5, 8)
            else:
                spacing = max(profile.bullet_spacing_after_pt, 2.5)

            paragraph.paragraph_format.space_after = Pt(spacing)

            add_run(
                paragraph,
                "• ",
                bold=True,
                size=Pt(profile.body_font_pt),
                color=BLUE,
            )
            add_run(
                paragraph,
                bullet.text,
                size=Pt(profile.body_font_pt),
            )


def render_resume(
    resume: Resume,
    output_path: str | Path,
    profile: LayoutProfile = NORMAL_PROFILE,
) -> Path:
    """
    Render a structured Resume into a deterministic DOCX file.

    Professional Experience intentionally begins on page 2.
    The selected LayoutProfile controls density only; resume content
    is never modified by the renderer.
    """

    output_path = Path(output_path)

    document = Document()
    configure_document(document, profile)

    add_header(document, resume)
    add_headline(document, resume.headline)
    add_summary(document, resume.summary, profile)
    add_skills(document, resume, profile)
    add_education(document, resume, profile)
    add_certifications(document, resume, profile)

    document.add_page_break()

    add_experience(document, resume, profile)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document.save(output_path)

    return output_path
