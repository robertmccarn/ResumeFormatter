from pathlib import Path

from docx import Document

from resume_formatter.models import (
    Bullet,
    Certification,
    Contact,
    Education,
    Experience,
    Resume,
    SkillCategory,
)
from resume_formatter.optimizer import COMPACT_PROFILE, NORMAL_PROFILE
from resume_formatter.renderer import render_resume


def make_resume() -> Resume:
    return Resume(
        contact=Contact(
            name="Robert McCarn",
            email="robert@example.com",
            phone="555-555-5555",
            location="Spring, TX",
            links=["linkedin.com/in/robert"],
        ),
        headline="Data Engineer",
        summary=(
            "Data Engineer with 4+ years of experience "
            "building and maintaining data solutions."
        ),
        skills=[
            SkillCategory(
                name="Databases",
                skills=["SQL Server", "Azure SQL"],
            ),
            SkillCategory(
                name="Languages",
                skills=["T-SQL", "Python", "PowerShell"],
            ),
        ],
        education=[
            Education(
                institution="University of Houston",
                degree="BBA, Management Information Systems",
                dates="2021",
            ),
        ],
        certifications=[
            Certification(
                name="Microsoft Fabric DP-700",
                date="2026",
            ),
        ],
        experience=[
            Experience(
                company="TEKsystems Global Services",
                title="Data Engineer",
                dates="2022–Present",
                bullets=[
                    Bullet(
                        text="Developed T-SQL stored procedures."
                    ),
                    Bullet(
                        text="Supported data modernization initiatives."
                    ),
                ],
            ),
        ],
    )


def test_render_resume_creates_docx(tmp_path: Path):
    output = tmp_path / "resume.docx"

    result = render_resume(make_resume(), output)

    assert result == output
    assert output.exists()
    assert output.stat().st_size > 0


def test_render_resume_contains_expected_content(tmp_path: Path):
    output = tmp_path / "resume.docx"

    render_resume(make_resume(), output)

    document = Document(output)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Robert McCarn" in text
    assert "Data Engineer" in text
    assert "SUMMARY" in text
    assert "Databases:" in text
    assert "University of Houston" in text
    assert "Microsoft Fabric DP-700" in text
    assert "PROFESSIONAL EXPERIENCE" in text
    assert "TEKsystems Global Services" in text


def test_render_resume_uses_bullet_markers_for_experience(
    tmp_path: Path,
):
    output = tmp_path / "resume.docx"

    render_resume(make_resume(), output)

    document = Document(output)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert "• Developed T-SQL stored procedures." in paragraphs
    assert "• Supported data modernization initiatives." in paragraphs


def test_render_resume_has_page_break_before_experience(
    tmp_path: Path,
):
    output = tmp_path / "resume.docx"

    render_resume(make_resume(), output)

    document = Document(output)

    body_xml = document._element.body.xml

    page_break_position = body_xml.find(
        'w:type="page"'
    )

    experience_position = body_xml.find(
        "PROFESSIONAL EXPERIENCE"
    )

    assert page_break_position != -1
    assert experience_position != -1

    assert page_break_position < experience_position


def test_render_resume_uses_selected_profile(tmp_path: Path):
    output = tmp_path / "compact.docx"

    render_resume(
        make_resume(),
        output,
        profile=COMPACT_PROFILE,
    )

    document = Document(output)

    normal = document.styles["Normal"]

    assert normal.font.size.pt == COMPACT_PROFILE.body_font_pt
    assert (
        normal.paragraph_format.space_after.pt
        == COMPACT_PROFILE.body_spacing_after_pt
    )


def test_render_resume_normal_profile_preserves_existing_defaults(
    tmp_path: Path,
):
    output = tmp_path / "normal.docx"

    render_resume(
        make_resume(),
        output,
        profile=NORMAL_PROFILE,
    )

    document = Document(output)

    normal = document.styles["Normal"]

    assert normal.font.size.pt == NORMAL_PROFILE.body_font_pt
    assert (
        normal.paragraph_format.space_after.pt
        == NORMAL_PROFILE.body_spacing_after_pt
    )


def test_render_profiles_preserve_resume_text(tmp_path: Path):
    normal_output = tmp_path / "normal.docx"
    compact_output = tmp_path / "compact.docx"

    resume = make_resume()

    render_resume(
        resume,
        normal_output,
        profile=NORMAL_PROFILE,
    )
    render_resume(
        resume,
        compact_output,
        profile=COMPACT_PROFILE,
    )

    normal_document = Document(normal_output)
    compact_document = Document(compact_output)

    normal_text = [
        paragraph.text
        for paragraph in normal_document.paragraphs
    ]
    compact_text = [
        paragraph.text
        for paragraph in compact_document.paragraphs
    ]

    assert normal_text == compact_text


def test_render_profiles_keep_experience_on_page_two(
    tmp_path: Path,
):
    output = tmp_path / "compact.docx"

    render_resume(
        make_resume(),
        output,
        profile=COMPACT_PROFILE,
    )

    document = Document(output)
    body_xml = document._element.body.xml

    page_break_position = body_xml.find(
        'w:type="page"'
    )
    experience_position = body_xml.find(
        "PROFESSIONAL EXPERIENCE"
    )

    assert page_break_position != -1
    assert experience_position != -1
    assert page_break_position < experience_position
